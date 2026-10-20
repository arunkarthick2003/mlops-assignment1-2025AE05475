"""Build the ten-section assignment report as a formatted Word document."""

from __future__ import annotations

import json
from pathlib import Path

import joblib
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sklearn.metrics import confusion_matrix
from sklearn.model_selection import train_test_split

from heart_mlops.model import validate_features

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
SCREENSHOTS = ROOT / "screenshots"
OUTPUT = ROOT / "Heart_Disease_MLOps_Report.docx"
REPOSITORY = "https://github.com/arunkarthick2003/mlops-assignment1-a3095887"
CI_RUN = f"{REPOSITORY}/actions/runs/39411595260"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def add_paragraph(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        add_paragraph(document, item, "List Bullet")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        set_cell_shading(cell, "17324D")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = str(value)
            if row_index % 2 == 0:
                set_cell_shading(cells[column_index], "EEF4F7")
            for paragraph in cells[column_index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        if widths:
            for column_index, width in enumerate(widths):
                cells[column_index].width = Inches(width)
    if widths:
        for column_index, width in enumerate(widths):
            table.rows[0].cells[column_index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(document: Document, path: Path, width: float = 6.8, max_height: float = 3.2) -> None:
    if not path.is_file():
        add_paragraph(document, f"Evidence image not found at build time: {path.relative_to(ROOT)}")
        return
    from PIL import Image

    with Image.open(path) as image:
        aspect = image.width / image.height
    image_width = min(width, max_height * aspect)
    image_height = image_width / aspect
    document.add_picture(str(path), width=Inches(image_width), height=Inches(image_height))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def make_architecture_image(path: Path) -> None:
    figure, axis = plt.subplots(figsize=(12, 4.2), facecolor="white")
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 4.2)
    axis.axis("off")
    steps = [
        (0.3, "UCI data", "download + clean"),
        (2.7, "Train", "CV + compare"),
        (5.1, "MLflow", "params + metrics"),
        (7.5, "Docker", "FastAPI image"),
        (9.9, "Kubernetes", "2 replicas"),
    ]
    for x_position, title, subtitle in steps:
        axis.text(
            x_position + 0.85,
            2.45,
            f"{title}\n{subtitle}",
            ha="center",
            va="center",
            fontsize=11,
            color="#17324D",
            bbox={"boxstyle": "round,pad=0.65", "facecolor": "#EEF4F7", "edgecolor": "#86A4B8", "linewidth": 1.4},
        )
    for x_position in (2.15, 4.55, 6.95, 9.35):
        axis.annotate("", (x_position + 0.35, 2.45), (x_position - 0.05, 2.45), arrowprops={"arrowstyle": "->", "color": "#E76F51", "lw": 2})
    axis.text(9.0, 0.65, "FastAPI: /health  /predict  /metrics", ha="center", fontsize=11, color="#287271")
    axis.annotate("", (9.85, 1.7), (9.25, 0.95), arrowprops={"arrowstyle": "->", "color": "#287271", "lw": 1.7})
    axis.text(3.2, 0.65, "GitHub Actions: Ruff -> train -> pytest -> EDA -> artifacts", ha="center", fontsize=10, color="#496174")
    axis.annotate("", (3.1, 1.55), (3.1, 0.95), arrowprops={"arrowstyle": "->", "color": "#496174", "lw": 1.5})
    figure.tight_layout()
    figure.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(figure)


def main() -> None:
    summary = json.loads((ARTIFACTS / "training_summary.json").read_text(encoding="utf-8"))
    frame = pd.read_csv(ROOT / "data/processed/heart_disease.csv")
    features = validate_features(frame)
    labels = frame["target"]
    _, x_test, _, y_test = train_test_split(
        features, labels, test_size=0.2, random_state=42, stratify=labels
    )
    model = joblib.load(ARTIFACTS / "heart_disease_model.joblib")
    predictions = model.predict(x_test)
    matrix = confusion_matrix(y_test, predictions, labels=[0, 1])
    architecture_path = ARTIFACTS / "architecture.png"
    make_architecture_image(architecture_path)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(39, 54, 66)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04
    for style_name, size, color in (("Title", 27, "17324D"), ("Heading 1", 18, "17324D"), ("Heading 2", 12, "287271")):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "AIMLCZG523 | Assignment 01 | Heart Disease MLOps"
    header.style = document.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("UCI Cleveland educational project  |  Page ")
    add_page_field(footer)

    # Page 1: cover and executive summary.
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("End-to-End Heart Disease\nMLOps Project")
    subtitle = add_paragraph(document, "Machine Learning Operations - Assignment 01")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(40, 114, 113)
    repo = add_paragraph(document, REPOSITORY)
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repo.runs[0].font.size = Pt(9)
    document.add_heading("Executive summary", level=1)
    add_paragraph(document, "This project builds a reproducible binary classifier for the Cleveland subset of the UCI Heart Disease dataset and carries it through data preparation, model comparison, MLflow tracking, automated checks, Docker, local Kubernetes, and an observable FastAPI service. It is an educational prototype, not clinical software.")
    selected = summary["selected_model"].replace("_", " ").title()
    chosen = next(run for run in summary["runs"] if run["model"] == summary["selected_model"])
    add_table(document, ["Selected classifier", "CV ROC-AUC", "Accuracy", "Precision", "Recall", "Holdout ROC-AUC"], [[selected + " (C=0.1)", f"{chosen['cv_roc_auc']:.4f}", f"{chosen['test_accuracy']:.4f}", f"{chosen['test_precision']:.4f}", f"{chosen['test_recall']:.4f}", f"{chosen['test_roc_auc']:.4f}"]], [1.6, 0.9, 0.9, 0.9, 0.8, 1.25])
    add_paragraph(document, "Verified delivery: six local tests passed; Ruff passed; the Docker image built; API health, prediction, and metrics were exercised; local Kubernetes reported two ready replicas; and the GitHub Actions run for commit 1dcd8b9 succeeded.")
    add_paragraph(document, "Safety note: predictions and confidence are not medical advice or validated clinical risk estimates.")
    document.add_page_break()

    # Page 2: problem, dataset, architecture.
    document.add_heading("1. Problem, dataset, and system architecture", level=1)
    add_paragraph(document, "The assignment asks for an end-to-end machine-learning solution that is reproducible, tracked, tested, containerized, deployable, and observable. The UCI Heart Disease Cleveland subset has 303 patient records and 13 predictors. Its source target `num` is integer-valued from 0 to 4; this project treats zero as absence and any positive value as disease present.")
    add_paragraph(document, "The architecture separates acquisition and EDA scripts, reusable sklearn preprocessing/model code, experiment tracking, API serving, automated CI, and deployment resources. The selected preprocessing/model pipeline is serialized once and copied into the serving image.")
    add_image(document, architecture_path, width=6.85, max_height=3.0)
    add_paragraph(document, "The service provides `/health`, `/predict`, and `/metrics`. The validated deployment target was Docker Desktop Kubernetes running locally; the LoadBalancer address is cluster-local, not a public internet endpoint.")
    add_paragraph(document, "Source: UCI Machine Learning Repository, Heart Disease dataset, DOI 10.24432/C52P4X.")
    document.add_page_break()

    # Page 3: acquisition and EDA.
    document.add_heading("2. Data acquisition and exploratory analysis", level=1)
    add_paragraph(document, "The download script retrieves the official UCI ZIP and extracts the processed Cleveland data file. Unknown cells represented by `?` are parsed as missing. The cleaned CSV contains 303 rows, 13 predictors, and a binary target. The class balance is 164 negative records and 139 positive records. Four `ca` values and two `thal` values are missing; all other predictor and target values are present.")
    add_paragraph(document, "The distribution figure summarizes age by target and class counts. The correlation heatmap is computed over numeric codes. The largest absolute Pearson associations with the target are thal (0.526), ca (0.460), exang (0.432), oldpeak (0.425), and thalach (-0.417). Categorical codes are nominal categories; those coefficients are descriptive and must not be read as causal effects.")
    add_image(document, ROOT / "artifacts/eda/distributions.png", width=6.6, max_height=2.65)
    add_image(document, ROOT / "artifacts/eda/correlation.png", width=5.4, max_height=3.7)
    add_paragraph(document, "The dataset is small and historical. EDA identifies patterns worth exploring but cannot establish population representativeness or clinical meaning.")
    document.add_page_break()

    # Page 4: engineering and protocol.
    document.add_heading("3. Feature engineering and training protocol", level=1)
    add_paragraph(document, "Numeric inputs are age, resting blood pressure, cholesterol, maximum heart rate, and ST depression. They receive median imputation and standard scaling. Categorical inputs are sex, chest-pain type, fasting blood sugar flag, resting ECG, exercise-induced angina, ST slope, vessel count, and thal category. They receive most-frequent imputation and one-hot encoding with unknown categories ignored.")
    add_paragraph(document, "All transformers are inside a scikit-learn ColumnTransformer/Pipeline, so they are fitted only on the training fold during cross-validation. This prevents preprocessing statistics from leaking from validation folds. The pipeline is reused unchanged for inference.")
    document.add_heading("Model selection", level=2)
    add_paragraph(document, "A stratified 80/20 split with random seed 42 creates 242 training rows and 61 holdout rows. Stratified five-fold CV on the training rows selects hyperparameters by ROC-AUC. Logistic Regression tests C = 0.1, 1, and 10 with balanced class weights. Random Forest tests max_depth in {None, 5} and min_samples_leaf in {1, 3}, also with balanced class weights and a fixed seed.")
    add_paragraph(document, "The holdout is reserved for final comparison. Accuracy, precision, recall, and ROC-AUC are reported, with disease present as the positive label. The small holdout means individual observations have a large effect on the results.")
    add_table(document, ["Stage", "Rows", "Purpose"], [["Training partition", "242", "5-fold stratified CV and grid search"], ["Holdout", "61", "Final evaluation after model selection"]], [1.7, 0.8, 4.4])
    document.add_page_break()

    # Page 5: results and errors.
    document.add_heading("4. Results and error analysis", level=1)
    rows = []
    for run in summary["runs"]:
        model_label = run["model"].replace("_", " ").title()
        params = ", ".join(f"{key.split('__')[-1]}={value}" for key, value in run["best_params"].items())
        rows.append([model_label, f"{run['cv_roc_auc']:.4f}", f"{run['test_accuracy']:.4f}", f"{run['test_precision']:.4f}", f"{run['test_recall']:.4f}", f"{run['test_roc_auc']:.4f}", params])
    add_table(document, ["Model", "CV AUC", "Accuracy", "Precision", "Recall", "Test AUC", "Best parameters"], rows, [1.2, 0.72, 0.82, 0.82, 0.72, 0.78, 1.8])
    add_paragraph(document, f"Selected model: {selected}. CV ROC-AUC differs by less than 0.001 between candidates. Logistic Regression has higher holdout accuracy, precision, and ROC-AUC; both have the same positive-class recall.")
    add_table(document, ["Actual \\ Predicted", "Absent (0)", "Present (1)"], [["Absent (0)", str(int(matrix[0, 0])), str(int(matrix[0, 1]))], ["Present (1)", str(int(matrix[1, 0])), str(int(matrix[1, 1]))]], [2.2, 1.4, 1.4])
    add_paragraph(document, "The holdout contains 28 true negatives, 5 false positives, 2 false negatives, and 26 true positives. This yields positive-class recall 26/28 = 0.9286. No threshold optimization or probability calibration was performed; the output confidence should not be assumed to be calibrated.")
    add_image(document, ARTIFACTS / "logistic_regression_evaluation.png", width=6.6, max_height=2.5)
    add_image(document, ARTIFACTS / "random_forest_evaluation.png", width=6.6, max_height=2.5)
    document.add_page_break()

    # Page 6: experiments and reproducibility.
    document.add_heading("5. Experiment tracking and reproducibility", level=1)
    add_paragraph(document, "MLflow experiment `uci-heart-disease-cleveland` is stored locally in `mlflow.db`. Model candidate runs log model identity, search parameters, seed, fold count, CV and holdout metrics, classification reports, and evaluation plots. A selected-model-package run logs the serialized pipeline and training summary as artifacts.")
    add_table(document, ["Candidate", "MLflow run ID", "CV ROC-AUC", "Holdout ROC-AUC"], [["Logistic Regression", "83dcf44e9edf4296ac210cd0c43fcbbe", "0.9028", "0.9654"], ["Random Forest", "5ba36aeea1d64c03a54d813b8a839b89", "0.9019", "0.9589"]], [1.5, 2.8, 1.2, 1.3])
    add_paragraph(document, "The final preprocessing/model pipeline is `artifacts/heart_disease_model.joblib`; `artifacts/training_summary.json` records selected model, features, metrics, parameters, and run IDs. A clean setup recreates raw data via the downloader, then EDA, training, and local experiment logging via the documented scripts.")
    add_image(document, SCREENSHOTS / "mlflow-runs.png", width=6.7, max_height=2.5)
    add_paragraph(document, "The figure above is a formatted metrics export built from the recorded run summary, not a raw MLflow browser screenshot. MLflow's local UI can be opened with `mlflow ui --backend-store-uri sqlite:///mlflow.db`.")
    document.add_page_break()

    # Page 7: API and Docker.
    document.add_heading("6. Model packaging, API, and Docker", level=1)
    add_paragraph(document, "The serialized object includes both preprocessing and the selected classifier, ensuring that serving uses the same transformations as training. The FastAPI request model validates all 13 required numeric fields and rejects extra fields. `/predict` returns the binary class, class label, confidence for that predicted class, and a clear educational-use disclaimer.")
    add_table(document, ["Endpoint", "Purpose", "Verification"], [["GET /health", "Model load status", "HTTP 200; model loaded"], ["POST /predict", "Single-record inference", "HTTP 200; JSON class/confidence"], ["GET /metrics", "Prometheus exposition", "HTTP 200; request counters/histograms"]], [1.35, 2.7, 2.7])
    add_paragraph(document, "The image is built as `heart-risk-api:local` from `python:3.12-slim`. The application runs as a non-root user on port 8000. The Docker image build completed successfully; the container's health, prediction, and metrics endpoints were exercised. Request logs include only request ID, method, path, status, and latency; feature payloads are not written to logs.")
    add_image(document, SCREENSHOTS / "docker-build.png", width=5.5, max_height=2.8)
    add_image(document, SCREENSHOTS / "api-predict.png", width=6.2, max_height=3.0)
    document.add_page_break()

    # Page 8: tests and CI.
    document.add_heading("7. Automated testing and CI/CD", level=1)
    add_paragraph(document, "The test suite includes four data/model pipeline tests and two API contract tests. It checks question-mark missing values, binary target transformation, required-feature validation, model fit/predict with missing and categorical values, successful health/prediction/metrics, and HTTP 422 for an incomplete payload. Six tests passed locally; Ruff checks passed.")
    add_paragraph(document, "GitHub Actions runs on push, pull request, and manual dispatch. The workflow checks out the source, installs Python 3.12 dependencies and the package, runs Ruff, trains and logs models, runs pytest, generates EDA figures, and uploads artifacts plus `mlflow.db`. Docker image construction is verified locally and is not a step in the current CI workflow.")
    add_table(document, ["CI run", "Commit", "Result", "Duration"], [["39411595260", "1dcd8b9", "Success", "1m 27s"]], [1.7, 1.5, 1.4, 1.4])
    add_paragraph(document, f"Repository: {REPOSITORY}\nWorkflow run: {CI_RUN}")
    add_image(document, SCREENSHOTS / "github-actions.png", width=6.7, max_height=3.3)
    document.add_page_break()

    # Page 9: deployment and monitoring.
    document.add_heading("8. Local Kubernetes deployment and monitoring", level=1)
    add_paragraph(document, "The deployment target is local Kubernetes through Docker Desktop. The manifest defines two API replicas, resource requests/limits, HTTP readiness and liveness probes, and a LoadBalancer Service. Both replicas became ready, and `kubectl rollout status deployment/heart-risk-api` reported a successful rollout. The local service was accessed with `kubectl port-forward service/heart-risk-api 8080:80`; `/health` returned `status: ok` and `model: loaded`.")
    add_paragraph(document, "The LoadBalancer address belongs to the local cluster and is not a public endpoint. The public deliverable is the source repository; runtime access instructions are local port-forwarding commands. No public cloud deployment is claimed.")
    add_image(document, SCREENSHOTS / "kubernetes-rollout.png", width=6.7, max_height=2.5)
    add_paragraph(document, "Monitoring instrumentation exposes request counts by method/path/status and latency histograms. Prometheus scrape configuration is supplied in `deployment/prometheus.yml`; API metrics and request logs were checked locally. Prometheus and Grafana were not deployed as a dashboard stack, so the demonstrated monitoring is the `/metrics` endpoint and application logs.")
    add_image(document, SCREENSHOTS / "api-metrics.png", width=6.5, max_height=2.8)
    document.add_page_break()

    # Page 10: limitations, conclusion, references.
    document.add_heading("9. Limitations, responsible use, and conclusion", level=1)
    add_paragraph(document, "The Cleveland data is small and historical. Only 61 observations were held out for final evaluation, so the confusion matrix and metrics have high sampling uncertainty. No external validation, calibration, prospective evaluation, threshold tuning, demographic subgroup analysis, or distribution-shift testing was conducted. Imputation and categorical encoding improve technical consistency but cannot remove sampling bias or establish clinical validity.")
    add_paragraph(document, "The API is a coursework artifact. Its prediction confidence is the model's class probability, not a calibrated clinical risk. The classifier must not be used for screening, diagnosis, treatment, or decisions about real patients. A credible research follow-up would require governed and representative datasets, independent validation, calibration and subgroup testing, model/data versioning, privacy and security review, monitoring alerts, and review by qualified medical professionals.")
    add_paragraph(document, "This project demonstrates the requested MLOps chain: reproducible data preparation, fold-safe preprocessing, two-candidate comparison, MLflow tracking, tests and CI, model packaging, Docker serving, local Kubernetes deployment, and basic request monitoring. The limitations above define what remains outside the assignment implementation.")
    document.add_heading("References", level=2)
    add_bullets(document, [
        "UCI Machine Learning Repository. Heart Disease dataset. DOI: 10.24432/C52P4X.",
        "scikit-learn documentation: Pipeline, ColumnTransformer, model selection, and classification metrics.",
        "MLflow documentation: Tracking experiments and backend stores.",
        "FastAPI documentation: Request models and API serving.",
        "Prometheus documentation: Client exposition format and scraping.",
    ])
    add_paragraph(document, "Submission note: a short pipeline walkthrough video is still to be recorded and added by the student.")

    document.core_properties.title = "End-to-End Heart Disease MLOps Report"
    document.core_properties.subject = "MLOps Assignment 01"
    document.core_properties.author = "arunkarthick2003"
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()